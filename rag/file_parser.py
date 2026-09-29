"""Bounded, non-executing text extraction, run in a disposable worker process."""

import csv
import io
import json
import sys
import zipfile
from pathlib import Path

from defusedxml import ElementTree

EXTENSIONS = (".pdf", ".md", ".xlsx", ".docx", ".csv", ".txt")
MAX_ARCHIVE_BYTES = 40 * 1024 * 1024
MAX_PAGES = 100
MAX_SHEETS = 50
MAX_ROWS = 10000
MAX_COLUMNS = 200
MAX_CELLS = 100000


class ParseError(Exception):
    """A safe, user-facing extraction failure."""

    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


class TextBuilder:
    """Reject expansion beyond the document budget; never silently truncate."""

    def __init__(self, limit: int):
        self.limit = limit
        self.parts: list[str] = []
        self.length = 0

    def add(self, text: str) -> None:
        if not text.strip():
            return
        self.length += len(text) + (2 if self.parts else 0)
        if self.length > self.limit:
            raise ParseError(
                "extracted_limit", "Extracted text exceeds the character limit. Split the file."
            )
        self.parts.append(text)

    def finish(self) -> str:
        text = "\n\n".join(self.parts)
        if not text.strip():
            raise ParseError(
                "empty_text", "No readable text found. Scanned PDFs need OCR before upload."
            )
        return text


def check_archive(path: Path, required: str) -> None:
    """Check Office ZIP expansion and reject encrypted/macro archives without extracting."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_ARCHIVE_BYTES:
            raise ParseError("archive_limit", "Office archive expands beyond the safety limit.")
        names = {entry.filename for entry in entries}
        if required not in names or "[Content_Types].xml" not in names:
            raise ParseError("corrupt_file", "The file is not a valid document of this format.")
        for entry in entries:
            if entry.flag_bits & 1 or "vbaproject" in entry.filename.lower():
                raise ParseError(
                    "encrypted_file", "Encrypted or macro-enabled Office files are not supported."
                )
            if (
                entry.file_size > 1024 * 1024
                and entry.file_size > max(1, entry.compress_size) * 300
            ):
                raise ParseError(
                    "archive_limit", "Office archive compression ratio exceeds the safety limit."
                )
            # Preflight XML with entity-safe parsing before downstream readers see it.
            if entry.filename.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(entry), forbid_dtd=True)


def decode_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        text = data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
    except UnicodeError as exc:
        raise ParseError(
            "encoding", "Save text files as UTF-8 or UTF-16 with a byte-order mark."
        ) from exc
    if any(ord(char) < 32 and char not in "\t\n\r" for char in text):
        raise ParseError(
            "binary_text", "This text file contains binary or unsupported control characters."
        )
    return text.replace("\r\n", "\n").replace("\r", "\n")


def extract(path: Path, extension: str, limit: int) -> dict:
    """Extract page/sheet/row-labelled text; no OCR, formula execution or remote fetching."""
    builder = TextBuilder(limit)
    warnings: list[str] = []
    if extension in {".txt", ".md"}:
        builder.add(decode_text(path))
    elif extension == ".csv":
        content = decode_text(path)
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        cells = 0
        for index, row in enumerate(csv.reader(io.StringIO(content), dialect, strict=True), 1):
            cells += len(row)
            if index > MAX_ROWS or len(row) > MAX_COLUMNS or cells > MAX_CELLS:
                raise ParseError(
                    "table_limit", "Table exceeds the row, column or cell limit. Split the file."
                )
            if any(value.strip() for value in row):
                builder.add(f"Row {index}: " + json.dumps(row, ensure_ascii=False))
    elif extension == ".docx":
        check_archive(path, "word/document.xml")
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"), forbid_dtd=True)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        body = root.find("w:body", ns)
        if body is None:
            raise ParseError("corrupt_file", "Word document has no body.")

        def paragraph(node) -> str:
            parts = []
            for item in node.iter():
                kind = item.tag.rsplit("}", 1)[-1]
                if kind == "t":
                    parts.append(item.text or "")
                elif kind in {"br", "cr", "tab"}:
                    parts.append("\t" if kind == "tab" else "\n")
            return "".join(parts)

        for block in body:
            if block.tag.endswith("}p"):
                builder.add(paragraph(block))
            elif block.tag.endswith("}tbl"):
                for index, row in enumerate(block.findall("w:tr", ns), 1):
                    values = [paragraph(cell) for cell in row.findall("w:tc", ns)]
                    builder.add(f"Table row {index}: " + json.dumps(values, ensure_ascii=False))
        warnings.append(
            "Word body paragraphs and tables extracted; images, headers and footers are not indexed."
        )
    elif extension == ".xlsx":
        check_archive(path, "xl/workbook.xml")
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = None
        try:
            cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            if len(workbook.worksheets) > MAX_SHEETS:
                raise ParseError("table_limit", "Workbook exceeds the sheet limit.")
            total_cells = 0
            missing_formula = False
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
                    raise ParseError("table_limit", "Sheet exceeds the row or column limit.")
                value_sheet = cached[sheet.title]
                # XML dimensions can be incorrect; reset and enforce limits while reading too.
                sheet.reset_dimensions()
                value_sheet.reset_dimensions()
                for row_index, (row, values) in enumerate(
                    zip(sheet.iter_rows(), value_sheet.iter_rows(), strict=True), 1
                ):
                    total_cells += len(row)
                    if row_index > MAX_ROWS or len(row) > MAX_COLUMNS or total_cells > MAX_CELLS:
                        raise ParseError(
                            "table_limit", "Workbook exceeds the row, column or cell limit."
                        )
                    output = []
                    for cell, value in zip(row, values, strict=True):
                        if cell.data_type == "f" and value.value is None:
                            missing_formula = True
                            output.append("[formula result unavailable]")
                        else:
                            output.append("" if value.value is None else str(value.value))
                    if any(output):
                        visibility = " (hidden)" if sheet.sheet_state != "visible" else ""
                        builder.add(
                            f"Sheet {sheet.title}{visibility}, row {row_index}: "
                            + json.dumps(output, ensure_ascii=False)
                        )
            if missing_formula:
                warnings.append(
                    "Some formulas have no cached values. Recalculate and save in Excel to index their results."
                )
            warnings.append(
                "All sheets, including hidden sheets, are indexed. Formulas are never executed."
            )
        finally:
            workbook.close()
            if cached is not None:
                cached.close()
    elif extension == ".pdf":
        if not path.read_bytes()[:1024].lstrip().startswith(b"%PDF-"):
            raise ParseError("corrupt_file", "The file is not a valid PDF.")
        import pdfplumber
        from pdfminer.pdfdocument import PDFPasswordIncorrect

        try:
            with pdfplumber.open(path) as pdf:
                if len(pdf.pages) > MAX_PAGES:
                    raise ParseError("page_limit", "PDF exceeds the page limit. Split the file.")
                for index, page in enumerate(pdf.pages, 1):
                    text = page.extract_text(x_tolerance_ratio=0.1) or ""
                    if text.strip():
                        builder.add(f"Page {index}\n{text}")
                    else:
                        warnings.append(
                            f"Page {index} has no extractable text; no OCR was performed."
                        )
                    for table_index, table in enumerate(
                        page.extract_tables(table_settings={"text_x_tolerance_ratio": 0.1}), 1
                    ):
                        for row_index, row in enumerate(table, 1):
                            if any(cell for cell in row):
                                builder.add(
                                    f"Page {index}, table {table_index}, row {row_index}: "
                                    + json.dumps(row, ensure_ascii=False)
                                )
                    page.close()
        except PDFPasswordIncorrect as exc:
            raise ParseError(
                "encrypted_file", "Password-protected PDFs must be decrypted before upload."
            ) from exc
        warnings.append(
            "PDF table extraction is best effort; table content may also appear in page text."
        )
    else:
        raise ParseError("unsupported_type", "Unsupported file extension.")
    return {"text": builder.finish(), "warnings": warnings}


def main() -> None:
    """Emit only a bounded result or a safe error; never expose parser internals."""
    try:
        result = extract(Path(sys.argv[1]), sys.argv[2], int(sys.argv[3]))
    except ParseError as exc:
        result = {"error": {"code": exc.code, "message": exc.message}}
    except Exception:
        result = {
            "error": {
                "code": "corrupt_file",
                "message": "Cannot parse this file. It may be corrupt, encrypted or incorrectly named.",
            }
        }
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
